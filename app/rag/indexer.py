
from __future__ import annotations

import hashlib
import sqlite3
import sys
from pathlib import Path

import numpy as np
from docx import Document
from openai import OpenAI
from pypdf import PdfReader

from app.rag.vault_db import get_vault_db, validate_vault

EMBED_MODEL = "text-embedding-nomic-embed-text-v1.5"
API_URL = "http://127.0.0.1:1234/v1"
BATCH_SIZE = 24
CHUNK_SIZE = 1400

client = OpenAI(base_url=API_URL, api_key="lm-studio", timeout=120)


def extract(path: Path) -> str:
    """Извлекает текст, не изменяя оригинальный документ."""
    suffix = path.suffix.lower()

    if suffix == ".md":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(parts)

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    return ""


def make_chunks(text: str) -> list[tuple[str, str]]:
    """Разбивает текст, стараясь сохранять заголовки Markdown."""
    chunks = []
    heading = ""
    buffer = []

    def flush():
        content = "\n".join(buffer).strip()
        if content:
            chunks.append((heading, content))
        buffer.clear()

    for line in text.splitlines():
        stripped = line.strip()

        if stripped.startswith("#"):
            flush()
            heading = stripped.lstrip("#").strip()
            continue

        if not stripped:
            if buffer:
                flush()
            continue

        current = "\n".join(buffer)
        if len(current) + len(stripped) + 1 > CHUNK_SIZE:
            flush()

        buffer.append(stripped)

        if len("\n".join(buffer)) > CHUNK_SIZE:
            flush()

    flush()

    result = []
    for title, content in chunks:
        for start in range(0, len(content), CHUNK_SIZE):
            part = content[start:start + CHUNK_SIZE].strip()
            if part:
                result.append((title, part))

    return result


def embed(texts: list[str]) -> list[list[float]]:
    response = client.embeddings.create(
        model=EMBED_MODEL,
        input=[f"search_document: {text}" for text in texts],
    )
    return [item.embedding for item in response.data]


def init_db(db: Path):
    db.parent.mkdir(parents=True, exist_ok=True)

    with sqlite3.connect(db) as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS documents (
                path TEXT PRIMARY KEY,
                sha256 TEXT NOT NULL,
                indexed_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        con.execute("""
            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                path TEXT NOT NULL,
                ordinal INTEGER NOT NULL,
                heading TEXT NOT NULL,
                content TEXT NOT NULL,
                embedding BLOB NOT NULL,
                dimensions INTEGER NOT NULL
            )
        """)
        con.execute("""
            CREATE INDEX IF NOT EXISTS idx_chunks_path
            ON chunks(path)
        """)


def files_to_index(root: Path) -> list[Path]:
    supported = {".md", ".docx", ".pdf"}
    excluded = {"node_modules", ".git"}

    return sorted(
        p for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in supported
        and not any(
            part.startswith(".")
            for part in p.relative_to(root).parts
        )
        and not any(
            part in excluded
            for part in p.relative_to(root).parts
        )
        and not p.name.startswith("~$")
    )


def index_file(
    path: Path,
    root: Path,
    con: sqlite3.Connection,
) -> str:
    relative = path.relative_to(root).as_posix()
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()

    row = con.execute(
        "SELECT sha256 FROM documents WHERE path = ?",
        (relative,),
    ).fetchone()

    if row and row[0] == digest:
        count = con.execute(
            "SELECT COUNT(*) FROM chunks WHERE path = ?",
            (relative,),
        ).fetchone()[0]

        if count > 0:
            return "без изменений"

    text = extract(path)
    parts = make_chunks(text)

    if not parts:
        with con:
            con.execute("DELETE FROM chunks WHERE path = ?", (relative,))
            con.execute("""
                INSERT OR REPLACE INTO documents(path, sha256)
                VALUES (?, ?)
            """, (relative, digest))
        return "пусто"

    new_rows = []

    for start in range(0, len(parts), BATCH_SIZE):
        batch = parts[start:start + BATCH_SIZE]

        vectors = embed([
            f"Документ: {relative}\n"
            f"Раздел: {heading or 'Без заголовка'}\n"
            f"Содержание:\n{content}"
            for heading, content in batch
        ])

        for offset, ((heading, content), vector) in enumerate(
            zip(batch, vectors)
        ):
            array = np.asarray(vector, dtype=np.float32)
            new_rows.append((
                relative,
                start + offset,
                heading,
                content,
                array.tobytes(),
                len(array),
            ))

    with con:
        con.execute("DELETE FROM chunks WHERE path = ?", (relative,))
        con.executemany("""
            INSERT INTO chunks(
                path, ordinal, heading, content, embedding, dimensions
            ) VALUES (?, ?, ?, ?, ?, ?)
        """, new_rows)
        con.execute("""
            INSERT OR REPLACE INTO documents(path, sha256)
            VALUES (?, ?)
        """, (relative, digest))

    return f"{len(new_rows)} фрагм."


def index_vault(vault_path: str) -> dict:
    root = validate_vault(vault_path)
    db = get_vault_db(str(root))

    init_db(db)
    paths = files_to_index(root)

    print(f"Хранилище: {root}")
    print(f"Документов: {len(paths)}")
    print(f"Индекс: {db}")

    indexed = 0
    errors = 0

    with sqlite3.connect(db) as con:
        for number, path in enumerate(paths, 1):
            try:
                status = index_file(path, root, con)
                print(f"[{number}/{len(paths)}] {path.name}: {status}")
                indexed += 1
            except Exception as exc:
                errors += 1
                print(
                    f"[ОШИБКА] {path}: "
                    f"{type(exc).__name__}: {exc}"
                )

        existing = {
            p.relative_to(root).as_posix()
            for p in paths
        }
        stored = {
            row[0]
            for row in con.execute("SELECT path FROM documents")
        }

        for relative in stored - existing:
            con.execute("DELETE FROM chunks WHERE path = ?", (relative,))
            con.execute("DELETE FROM documents WHERE path = ?", (relative,))

        con.commit()

    print(f"Индексация завершена. Успешно: {indexed}, ошибок: {errors}")

    return {
        "vault_path": str(root),
        "database": str(db),
        "documents": len(paths),
        "indexed": indexed,
        "errors": errors,
    }


def main():
    if len(sys.argv) < 2:
        raise SystemExit(
            "Использование: python -m app.rag.indexer /путь/к/Obsidian"
        )

    index_vault(sys.argv[1])


if __name__ == "__main__":
    main()
